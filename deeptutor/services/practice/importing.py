"""Bounded, non-executing parsers for the learner's CSV, XLSX and JSON files."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
import zipfile

from .normalization import cell_text, normalize_question

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 64


_cell = cell_text


def parse_rows(data: bytes, filename: str) -> list[tuple[int, dict]]:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Choose a non-empty file up to 5 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        rows = json.loads(data.decode("utf-8-sig"))
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError("JSON must contain an array of question objects")
        if len(rows) > MAX_ROWS:
            raise ValueError("Import at most 500 questions at a time")
        return list(enumerate(rows, 1))
    if suffix in {".csv", ".tsv"}:
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("gb18030")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        rows = []
        try:
            for row in csv.reader(io.StringIO(content, newline=""), dialect):
                if len(row) > MAX_COLUMNS:
                    raise ValueError("Import at most 64 columns at a time")
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("Import at most 500 questions at a time")
        except csv.Error as exc:
            raise ValueError("Invalid CSV data or a cell exceeds the supported size") from exc
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("The expanded workbook exceeds 30 MB")
            if len(archive.infolist()) > 2000:
                raise ValueError("The workbook contains too many parts")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.active
            if sheet is None:
                raise ValueError("The workbook has no active worksheet")
            if (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("Import at most 64 columns at a time")
            rows = []
            for cells in sheet.iter_rows():
                if len(cells) > MAX_COLUMNS:
                    raise ValueError("Import at most 64 columns at a time")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Replace spreadsheet formulas with values before importing")
                rows.append([cell.value for cell in cells])
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("Import at most 500 questions at a time")
        finally:
            book.close()
    else:
        raise ValueError("Supported formats: .xlsx, .csv, .tsv and .json")
    if not rows:
        raise ValueError("The file is empty")
    headers = [_cell(value).lower() for value in rows[0]]
    nonempty = [header for header in headers if header]
    if len(nonempty) != len(set(nonempty)):
        raise ValueError("Column names must be unique")
    if len(rows) > MAX_ROWS + 1:
        raise ValueError("Import at most 500 questions at a time")
    return [
        (n, dict(zip(headers, row, strict=False)))
        for n, row in enumerate(rows[1:], 2)
        if any(_cell(value) for value in row)
    ]


def preview(data: bytes, filename: str) -> dict:
    questions, errors = [], []
    for number, row in parse_rows(data, filename):
        try:
            questions.append(normalize_question(row))
        except (ValueError, TypeError) as exc:
            errors.append({"row": number, "message": str(exc)})
    if not questions and not errors:
        raise ValueError("The file contains no questions")
    return {"questions": questions, "errors": errors}
